"""
Генерація офіційного звіту з Лабораторної роботи № 1 у форматі DOCX (Microsoft Word).
Вимоги оформлення:
- Жодного жирного шрифту (font-weight normal усюди);
- Уникати списків (без маркерів • та дефісних переліків, суцільний та структурований текст);
- Жодних довгих тире (виключно звичайний дефіс -);
- Усі рисунки обов'язково підписані;
- Коефіцієнт k = 0.5, варіант 4 (N = 4);
- Формат А4, поля по 20 мм (2.0 см);
- Шрифт Times New Roman, 14 пт, міжрядковий інтервал 1.5, абзацний відступ 1.25 см;
- Рік: 2026.
"""

import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Mm, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="{kwargs.get("top", "single")}" w:sz="{kwargs.get("top_sz", "4")}" w:color="{kwargs.get("color", "000000")}"/>\n'
        f'  <w:left w:val="{kwargs.get("left", "single")}" w:sz="{kwargs.get("left_sz", "4")}" w:color="{kwargs.get("color", "000000")}"/>\n'
        f'  <w:bottom w:val="{kwargs.get("bottom", "single")}" w:sz="{kwargs.get("bottom_sz", "4")}" w:color="{kwargs.get("color", "000000")}"/>\n'
        f'  <w:right w:val="{kwargs.get("right", "single")}" w:sz="{kwargs.get("right_sz", "4")}" w:color="{kwargs.get("color", "000000")}"/>\n'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)

def set_cell_shading(cell, color_hex="F2F2F2"):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)

def create_report_docx():
    doc = Document()
    
    # Налаштування параметрів сторінки А4 та полів
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)
    section.left_margin = Mm(20)
    section.right_margin = Mm(20)
    
    # Стиль Normal: Times New Roman, 14 pt, 1.5 інтервал, без жирного
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(14)
    normal_style.font.bold = False
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    normal_style.paragraph_format.line_spacing = 1.5
    normal_style.paragraph_format.space_after = Pt(4)
    normal_style.paragraph_format.space_before = Pt(0)
    
    base_dir = os.path.dirname(os.path.abspath(__file__))
    assets_dir = os.path.join(base_dir, "..", "assets")
    output_docx = os.path.join(base_dir, "..", "Звіт_ЛР1_Буханцев.docx")
    
    # ==================== ТИТУЛЬНИЙ АРКУШ ====================
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.line_spacing = 1.15
    p_inst.paragraph_format.first_line_indent = Cm(0)
    p_inst.paragraph_format.space_after = Pt(24)
    
    run_inst = p_inst.add_run(
        "МІНІСТЕРСТВО ОСВІТИ І НАУКИ УКРАЇНИ\n"
        "КРЕМЕНЧУЦЬКИЙ НАЦІОНАЛЬНИЙ УНІВЕРСИТЕТ\n"
        "ІМЕНІ МИХАЙЛА ОСТРОГРАДСЬКОГО\n"
        "НАВЧАЛЬНО-НАУКОВИЙ ІНСТИТУТ ЕЛЕКТРИЧНОЇ ІНЖЕНЕРІЇ\n"
        "ТА ІНФОРМАЦІЙНИХ ТЕХНОЛОГІЙ\n"
        "КАФЕДРА АВТОМАТИЗАЦІЇ ТА ІНФОРМАЦІЙНИХ СИСТЕМ"
    )
    run_inst.bold = False
    run_inst.font.size = Pt(12)
    
    p_course = doc.add_paragraph()
    p_course.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_course.paragraph_format.line_spacing = 1.2
    p_course.paragraph_format.first_line_indent = Cm(0)
    p_course.paragraph_format.space_before = Pt(40)
    p_course.paragraph_format.space_after = Pt(40)
    
    r_sub = p_course.add_run("ОСВІТНІЙ КОМПОНЕНТ\n«ДОСЛІДЖЕННЯ ОПЕРАЦІЙ»\n\n")
    r_sub.bold = False
    r_sub.font.size = Pt(14)
    
    r_rep = p_course.add_run("ЗВІТ\n")
    r_rep.bold = False
    r_rep.font.size = Pt(18)
    
    r_lab = p_course.add_run("З ЛАБОРАТОРНОЇ РОБОТИ №1")
    r_lab.bold = False
    r_lab.font.size = Pt(15)
    
    # Блок виконавця та перевіряючого
    p_sign = doc.add_paragraph()
    p_sign.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sign.paragraph_format.line_spacing = 1.25
    p_sign.paragraph_format.first_line_indent = Cm(0)
    p_sign.paragraph_format.space_before = Pt(50)
    p_sign.paragraph_format.space_after = Pt(70)
    
    r_sign = p_sign.add_run(
        "Виконав:\n"
        "здобувач групи КН-24-1\n"
        "Буханцев М. В.\n\n"
        "Перевірила:\n"
        "доцент кафедри АІС\n"
        "Бурдільна Є. В."
    )
    r_sign.bold = False
    r_sign.font.size = Pt(13)
    
    # Місто і рік
    p_city = doc.add_paragraph()
    p_city.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_city.paragraph_format.first_line_indent = Cm(0)
    p_city.paragraph_format.space_after = Pt(0)
    r_city = p_city.add_run("Кременчук 2026")
    r_city.bold = False
    r_city.font.size = Pt(13)
    
    # Розрив сторінки після титулки
    doc.add_page_break()
    
    # ==================== ОСНОВНА ЧАСТИНА ====================
    # Тема і Мета
    p_meta1 = doc.add_paragraph()
    p_meta1.paragraph_format.first_line_indent = Cm(0)
    p_meta1.paragraph_format.space_before = Pt(6)
    p_meta1.paragraph_format.space_after = Pt(4)
    r = p_meta1.add_run("Тема: Двомірна задача лінійного програмування.")
    r.bold = False
    
    p_meta2 = doc.add_paragraph()
    p_meta2.paragraph_format.first_line_indent = Cm(0)
    p_meta2.paragraph_format.space_after = Pt(12)
    r = p_meta2.add_run("Мета: набути навичок з розв’язування двомірних задач лінійного програмування графічним методом.")
    r.bold = False
    
    # Хід роботи
    p_hid = doc.add_paragraph()
    p_hid.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_hid.paragraph_format.first_line_indent = Cm(0)
    p_hid.paragraph_format.space_before = Pt(8)
    p_hid.paragraph_format.space_after = Pt(10)
    r_hid = p_hid.add_run("Хід роботи")
    r_hid.bold = False
    r_hid.font.size = Pt(14)
    
    # Текст умови
    p_text1 = doc.add_paragraph()
    p_text1.paragraph_format.first_line_indent = Cm(1.25)
    p_text1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p_text1.add_run(
        "Підприємство випускає столи двох моделей: А і В. Для випуску одного столу моделі А потрібно SA одиниці сировини "
        "та TA одиниці машинного часу. Для випуску одного столу моделі В потрібно SB одиниці сировини та TB одиниць машинного часу. "
        "Прибуток від реалізації одного столу моделі А складає WA грошові одиниці, столу моделі В - WB грошові одиниці. "
        "На підприємстві наявні 1700 одиниць сировини та 1600 одиниць машинного часу. Визначити, яким має бути план виробництва, "
        "щоб підприємство отримало максимальний прибуток."
    )
    r.bold = False
    
    p_text2 = doc.add_paragraph()
    p_text2.paragraph_format.first_line_indent = Cm(1.25)
    p_text2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p_text2.add_run(
        "За індивідуальним завданням варіанта 4 (номер у списку N = 4, коефіцієнт підгрупи 1 k = 0.5) "
        "коефіцієнти розраховуються за формулами:"
    )
    r.bold = False
    
    # Розрахунок параметрів без маркованих списків
    params = [
        "SA = 0.5 · N · k = 0.5 · 4 · 0.5 = 1.0 (од. сировини на стіл А);",
        "TA = N / k = 4 / 0.5 = 8.0 (год. машинного часу на стіл А);",
        "SB = N = 4.0 (од. сировини на стіл В);",
        "TB = 0.4 · N · k = 0.4 · 4 · 0.5 = 0.8 (год. машинного часу на стіл В);",
        "WA = N · k = 4 · 0.5 = 2.0 (грн прибутку від 1 столу А);",
        "WB = 1.2 · N · k = 1.2 · 4 · 0.5 = 2.4 (грн прибутку від 1 столу В)."
    ]
    for param in params:
        p_p = doc.add_paragraph()
        p_p.paragraph_format.first_line_indent = Cm(1.25)
        p_p.paragraph_format.space_after = Pt(2)
        r = p_p.add_run(param)
        r.bold = False
    
    # Таблиця 1
    p_t1_title = doc.add_paragraph()
    p_t1_title.paragraph_format.first_line_indent = Cm(0)
    p_t1_title.paragraph_format.space_before = Pt(8)
    p_t1_title.paragraph_format.space_after = Pt(2)
    r = p_t1_title.add_run("Таблиця 1 - Вихідні дані задачі оптимізації виробництва (Варіант 4)")
    r.bold = False
    
    table1 = doc.add_table(rows=4, cols=4)
    table1.alignment = WD_TABLE_ALIGNMENT.CENTER
    table1.autofit = False
    
    t1_headers = ["Показник", "Стіл моделі А (x1)", "Стіл моделі В (x2)", "Наявний запас на складі"]
    t1_data = [
        ["Витрати сировини на виріб, од.", "1.0", "4.0", "1700 од."],
        ["Витрати машинного часу, год.", "8.0", "0.8", "1600 год."],
        ["Прибуток від реалізації виробу, грош. од.", "2.0", "2.4", "-"]
    ]
    
    for c_idx, h_text in enumerate(t1_headers):
        cell = table1.cell(0, c_idx)
        cell.text = h_text
        set_cell_shading(cell, "F2F2F2")
        set_cell_border(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].bold = False
        p.runs[0].font.size = Pt(12)
        
    for r_idx, row_values in enumerate(t1_data):
        for c_idx, val in enumerate(row_values):
            cell = table1.cell(r_idx + 1, c_idx)
            cell.text = val
            set_cell_border(cell)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.runs[0].bold = False
            p.runs[0].font.size = Pt(12)
            
    col_widths = [Cm(7.0), Cm(3.2), Cm(3.2), Cm(3.6)]
    for row in table1.rows:
        for c_idx, w in enumerate(col_widths):
            row.cells[c_idx].width = w

    # 2. Економіко-математична модель ЗЛП
    p_sec2 = doc.add_paragraph()
    p_sec2.paragraph_format.first_line_indent = Cm(1.25)
    p_sec2.paragraph_format.space_before = Pt(12)
    r = p_sec2.add_run("2. Економіко-математична модель ЗЛП")
    r.bold = False
    
    p_m1 = doc.add_paragraph()
    p_m1.paragraph_format.first_line_indent = Cm(1.25)
    p_m1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p_m1.add_run("Нехай x1 - щоденний обсяг випуску столів моделі А (шт.), а x2 - щоденний обсяг випуску столів моделі В (шт.).")
    r.bold = False
    
    p_m2 = doc.add_paragraph()
    p_m2.paragraph_format.first_line_indent = Cm(1.25)
    r = p_m2.add_run("Цільова функція (критерій ефективності):")
    r.bold = False
    
    p_cf = doc.add_paragraph()
    p_cf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cf.paragraph_format.first_line_indent = Cm(0)
    r = p_cf.add_run("W(x1, x2) = 2.0 · x1 + 2.4 · x2 → max")
    r.bold = False
    r.font.size = Pt(14)
    
    p_m3 = doc.add_paragraph()
    p_m3.paragraph_format.first_line_indent = Cm(1.25)
    r = p_m3.add_run("Система обмежень за виробничими ресурсами:")
    r.bold = False
    
    p_sys = doc.add_paragraph()
    p_sys.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sys.paragraph_format.first_line_indent = Cm(0)
    r = p_sys.add_run(
        "1.0 · x1 + 4.0 · x2 ≤ 1700  (обмеження за сировиною)\n"
        "8.0 · x1 + 0.8 · x2 ≤ 1600  (обмеження за машинним часом)"
    )
    r.bold = False
    
    p_m4 = doc.add_paragraph()
    p_m4.paragraph_format.first_line_indent = Cm(1.25)
    r = p_m4.add_run("Граничні умови невід'ємності змінних:")
    r.bold = False
    
    p_nn = doc.add_paragraph()
    p_nn.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_nn.paragraph_format.first_line_indent = Cm(0)
    r = p_nn.add_run("x1 ≥ 0,   x2 ≥ 0")
    r.bold = False

    # 3. Графічний розв'язок ЗЛП
    p_sec3 = doc.add_paragraph()
    p_sec3.paragraph_format.first_line_indent = Cm(1.25)
    p_sec3.paragraph_format.space_before = Pt(12)
    r = p_sec3.add_run("3. Графічний розв'язок ЗЛП")
    r.bold = False
    
    p_gr1 = doc.add_paragraph()
    p_gr1.paragraph_format.first_line_indent = Cm(1.25)
    p_gr1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p_gr1.add_run(
        "Перетворимо нерівності обмежень у рівності граничних прямих:\n"
        "1. Пряма L1 (сировина): 1.0 x1 + 4.0 x2 = 1700 ⇒ x2 = 425 - 0.25 x1. "
        "Точки перетину з осями координат: (0; 425) та (1700; 0).\n"
        "2. Пряма L2 (час): 8.0 x1 + 0.8 x2 = 1600 ⇒ x2 = 2000 - 10 x1. "
        "Точки перетину з осями координат: (0; 2000) та (200; 0)."
    )
    r.bold = False
    
    p_gr2 = doc.add_paragraph()
    p_gr2.paragraph_format.first_line_indent = Cm(1.25)
    p_gr2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p_gr2.add_run(
        "Перетин півплощин у першому координатному квадранті утворює область допустимих розв'язків (ОДР) "
        "у вигляді опуклого замкненого чотирикутника OABC з вершинами: "
        "вершина O(0; 0) - початок координат; "
        "вершина A(200; 0) - точка перетину прямої L2 з віссю Ox1; "
        "вершина B(x1*; x2*) - точка взаємного перетину прямих обмежень L1 та L2; "
        "вершина C(0; 425) - точка перетину прямої L1 з віссю Ox2."
    )
    r.bold = False
    
    p_gr3 = doc.add_paragraph()
    p_gr3.paragraph_format.first_line_indent = Cm(1.25)
    p_gr3.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p_gr3.add_run(
        "Знайдемо аналітично координати вершини B розв'язанням системи лінійних рівнянь:\n"
        "   1.0 x1 + 4.0 x2 = 1700\n"
        "   8.0 x1 + 0.8 x2 = 1600\n"
        "Помноживши перше рівняння на 8 та віднявши друге, отримуємо: 31.2 x2 = 12000 ⇒ x2* = 5000 / 13 ≈ 384.62.\n"
        "Підставивши у перше рівняння, маємо: 1.0 x1 = 1700 - 4(384.615) = 2100 / 13 ≈ 161.54."
    )
    r.bold = False
    
    p_gr4 = doc.add_paragraph()
    p_gr4.paragraph_format.first_line_indent = Cm(1.25)
    p_gr4.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p_gr4.add_run(
        "Вектор градієнта цільової функції: c = grad W = (2.0; 2.4). "
        "Він перпендикулярний до ліній рівня 2.0 x1 + 2.4 x2 = const. "
        "Переміщуючи лінію рівня у напрямку градієнта від початку координат, встановлюємо, "
        "що крайньою точкою виходу лінії рівня з багатокутника ОДР є вершина B(161.54; 384.62)."
    )
    r.bold = False
    
    # Рисунок 1 (з обов'язковим підписом)
    fig1_path = os.path.join(assets_dir, "graph_lab1.png")
    if os.path.exists(fig1_path):
        p_fig1 = doc.add_paragraph()
        p_fig1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_fig1.paragraph_format.first_line_indent = Cm(0)
        p_fig1.paragraph_format.space_before = Pt(8)
        p_fig1.paragraph_format.space_after = Pt(2)
        doc.add_picture(fig1_path, width=Inches(6.0))
        
        p_cap1 = doc.add_paragraph()
        p_cap1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap1.paragraph_format.first_line_indent = Cm(0)
        p_cap1.paragraph_format.space_after = Pt(12)
        r = p_cap1.add_run("Рисунок 1 - Графічне розв'язання задачі лінійного програмування (Варіант 4)")
        r.bold = False
        r.font.size = Pt(12)

    # 4. Розрахунок цільової функції у кутових точках
    p_sec4 = doc.add_paragraph()
    p_sec4.paragraph_format.first_line_indent = Cm(1.25)
    p_sec4.paragraph_format.space_before = Pt(10)
    r = p_sec4.add_run("4. Розрахунок цільової функції у кутових точках ОДР")
    r.bold = False
    
    p_pts_intro = doc.add_paragraph()
    p_pts_intro.paragraph_format.first_line_indent = Cm(1.25)
    r = p_pts_intro.add_run("Значення цільової функції у кутових вершинах багатокутника рішень становлять:")
    r.bold = False
    
    pts_vals = [
        "W(O) = 2.0 · 0 + 2.4 · 0 = 0.00 грош. од.;",
        "W(A) = 2.0 · 200 + 2.4 · 0 = 400.00 грош. од.;",
        "W(C) = 2.0 · 0 + 2.4 · 425 = 1020.00 грош. од.;",
        "W(B) = 2.0 · 161.54 + 2.4 · 384.62 = 323.08 + 923.08 = 1246.15 грош. од."
    ]
    for pt in pts_vals:
        p_p = doc.add_paragraph()
        p_p.paragraph_format.first_line_indent = Cm(1.25)
        p_p.paragraph_format.space_after = Pt(2)
        r = p_p.add_run(pt)
        r.bold = False
        
    p_wmax = doc.add_paragraph()
    p_wmax.paragraph_format.first_line_indent = Cm(1.25)
    r = p_wmax.add_run("Отже, максимальний прибуток досягається у точці B і становить W_max = 1246.15 грн.")
    r.bold = False

    # 5. Обчислення залишків на складах
    p_sec5 = doc.add_paragraph()
    p_sec5.paragraph_format.first_line_indent = Cm(1.25)
    p_sec5.paragraph_format.space_before = Pt(10)
    r = p_sec5.add_run("5. Обчислення залишків ресурсів на складах")
    r.bold = False
    
    p_left = doc.add_paragraph()
    p_left.paragraph_format.first_line_indent = Cm(1.25)
    p_left.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p_left.add_run(
        "Підставимо отримані оптимальні значення x1* = 161.54 та x2* = 384.62 у ліві частини обмежень:\n"
        "Використання сировини: 1.0 · 161.54 + 4.0 · 384.62 = 1700.0 од. Залишок: 1700 - 1700 = 0.0 од. (100% використання).\n"
        "Використання машинного часу: 8.0 · 161.54 + 0.8 · 384.62 = 1600.0 год. Залишок: 1600 - 1600 = 0.0 год. (100% використання).\n"
        "Висновок: обидва ресурси є повністю вичерпаними (дефіцитними), резервів немає."
    )
    r.bold = False

    # 6. Цілочисельний аналіз
    p_sec6 = doc.add_paragraph()
    p_sec6.paragraph_format.first_line_indent = Cm(1.25)
    p_sec6.paragraph_format.space_before = Pt(10)
    r = p_sec6.add_run("6. Цілочисельний аналіз для неподільних виробів")
    r.bold = False
    
    p_int1 = doc.add_paragraph()
    p_int1.paragraph_format.first_line_indent = Cm(1.25)
    p_int1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p_int1.add_run(
        "Оскільки столи є фізично неподільною продукцією (x1, x2 належать множині цілих чисел Z), "
        "проведено аналіз цілочисельних вузлів сітки в околі вершини B(161.54; 384.62) з перевіркою виконання обмежень."
    )
    r.bold = False
    
    # Рисунок 2 (з обов'язковим підписом)
    fig2_path = os.path.join(assets_dir, "graph_lab1_zoom.png")
    if os.path.exists(fig2_path):
        p_fig2 = doc.add_paragraph()
        p_fig2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_fig2.paragraph_format.first_line_indent = Cm(0)
        p_fig2.paragraph_format.space_before = Pt(8)
        p_fig2.paragraph_format.space_after = Pt(2)
        doc.add_picture(fig2_path, width=Inches(5.8))
        
        p_cap2 = doc.add_paragraph()
        p_cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap2.paragraph_format.first_line_indent = Cm(0)
        p_cap2.paragraph_format.space_after = Pt(10)
        r = p_cap2.add_run("Рисунок 2 - Деталізований окіл точки оптимуму та цілочисельний аналіз")
        r.bold = False
        r.font.size = Pt(12)
        
    p_int2 = doc.add_paragraph()
    p_int2.paragraph_format.first_line_indent = Cm(1.25)
    p_int2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p_int2.add_run(
        "Серед усіх допустимих цілочисельних планів найкращим є план (160; 385):\n"
        "Випуск: 160 столів моделі А та 385 столів моделі В;\n"
        "Прибуток: W = 2.0(160) + 2.4(385) = 320 + 924 = 1244.00 грн;\n"
        "Використання сировини: 1.0(160) + 4.0(385) = 1700 од. (залишок 0 од., 100% завантаження);\n"
        "Використання часу: 8.0(160) + 0.8(385) = 1588.0 год. (залишок 12.0 год., завантаження 99.25%)."
    )
    r.bold = False

    # 7. Зведена таблиця та рекомендації
    p_sec7 = doc.add_paragraph()
    p_sec7.paragraph_format.first_line_indent = Cm(1.25)
    p_sec7.paragraph_format.space_before = Pt(10)
    r = p_sec7.add_run("7. Зведена таблиця результатів та рекомендації виробництву")
    r.bold = False
    
    # Рисунок 3 (з обов'язковим підписом)
    fig3_path = os.path.join(assets_dir, "results_summary.png")
    if os.path.exists(fig3_path):
        p_fig3 = doc.add_paragraph()
        p_fig3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_fig3.paragraph_format.first_line_indent = Cm(0)
        p_fig3.paragraph_format.space_before = Pt(6)
        p_fig3.paragraph_format.space_after = Pt(2)
        doc.add_picture(fig3_path, width=Inches(6.0))
        
        p_cap3 = doc.add_paragraph()
        p_cap3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap3.paragraph_format.first_line_indent = Cm(0)
        p_cap3.paragraph_format.space_after = Pt(10)
        r = p_cap3.add_run("Рисунок 3 - Зведені показники розв'язку оптимізаційної задачі")
        r.bold = False
        r.font.size = Pt(12)
        
    p_t2_title = doc.add_paragraph()
    p_t2_title.paragraph_format.first_line_indent = Cm(0)
    p_t2_title.paragraph_format.space_before = Pt(8)
    p_t2_title.paragraph_format.space_after = Pt(2)
    r = p_t2_title.add_run("Таблиця 2 - Порівняльний аналіз отриманих планів виробництва")
    r.bold = False
    
    table2 = doc.add_table(rows=8, cols=3)
    table2.alignment = WD_TABLE_ALIGNMENT.CENTER
    table2.autofit = False
    
    t2_headers = ["Показник", "Неперервний розв'язок (x*)", "Практичний цілочисельний план"]
    t2_data = [
        ["Столи моделі А (x1), шт.", "161.54", "160"],
        ["Столи моделі В (x2), шт.", "384.62", "385"],
        ["Максимальний прибуток (W), грн", "1246.15", "1244.00"],
        ["Використання сировини", "1700 / 1700 од. (100%)", "1700 / 1700 од. (100%)"],
        ["Залишок сировини", "0.0 од.", "0.0 од."],
        ["Використання машинного часу", "1600 / 1600 год. (100%)", "1588.0 / 1600 год. (99.25%)"],
        ["Залишок машинного часу", "0.0 год.", "12.0 год."]
    ]
    
    for c_idx, h_text in enumerate(t2_headers):
        cell = table2.cell(0, c_idx)
        cell.text = h_text
        set_cell_shading(cell, "F2F2F2")
        set_cell_border(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].bold = False
        p.runs[0].font.size = Pt(12)
        
    for r_idx, row_values in enumerate(t2_data):
        for c_idx, val in enumerate(row_values):
            cell = table2.cell(r_idx + 1, c_idx)
            cell.text = val
            set_cell_border(cell)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.runs[0].bold = False
            p.runs[0].font.size = Pt(12)
            
    col_widths2 = [Cm(8.0), Cm(4.5), Cm(4.5)]
    for row in table2.rows:
        for c_idx, w in enumerate(col_widths2):
            row.cells[c_idx].width = w
            
    p_rec = doc.add_paragraph()
    p_rec.paragraph_format.first_line_indent = Cm(1.25)
    p_rec.paragraph_format.space_before = Pt(8)
    p_rec.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p_rec.add_run(
        "Рекомендації щодо оптимізації виробництва:\n"
        "Оскільки ресурс сировини вичерпаний на 100%, а машинний час задіяний на 99.25%, підприємству слід "
        "затвердити щоденний виробничий план у розмірі 160 столів моделі А та 385 столів моделі В. "
        "Це забезпечує отримання максимального практичного прибутку у сумі 1244.00 грн за мінімального недовантаження "
        "обладнання (лише 12 годин за зміну). Подальше збільшення прибутку можливе виключно за умови додаткового постачання сировини."
    )
    r.bold = False

    # 8. Програмна реалізація мовою Python
    p_sec8 = doc.add_paragraph()
    p_sec8.paragraph_format.first_line_indent = Cm(1.25)
    p_sec8.paragraph_format.space_before = Pt(12)
    r = p_sec8.add_run("8. Програмна реалізація мовою Python (лістинг розрахунку)")
    r.bold = False
    
    code_text = (
        "import numpy as np\n"
        "from scipy.optimize import linprog\n\n"
        "N, k = 4, 0.5\n"
        "S_A, T_A, W_A = 0.5 * N * k, N / k, N * k             # 1.0, 8.0, 2.0\n"
        "S_B, T_B, W_B = float(N), 0.4 * N * k, 1.2 * N * k       # 4.0, 0.8, 2.4\n"
        "LIMIT_RAW, LIMIT_TIME = 1700.0, 1600.0\n\n"
        "# Чисельний розв'язок за допомогою SciPy HiGHS\n"
        "c = [-W_A, -W_B]\n"
        "A_ub = [[S_A, S_B], [T_A, T_B]]\n"
        "b_ub = [LIMIT_RAW, LIMIT_TIME]\n"
        "bounds = [(0, None), (0, None)]\n\n"
        "res_cont = linprog(c, A_ub=A_ub, b_ub=b_ub, bounds=bounds, method='highs')\n"
        "res_int = linprog(c, A_ub=A_ub, b_ub=b_ub, bounds=bounds, integrality=[1, 1], method='highs')\n\n"
        "print(f'Неперервний: x1={res_cont.x[0]:.2f}, x2={res_cont.x[1]:.2f}, W={-res_cont.fun:.2f}')\n"
        "print(f'Цілочисельний: x1={res_int.x[0]:.0f}, x2={res_int.x[1]:.0f}, W={-res_int.fun:.2f}')"
    )
    p_code = doc.add_paragraph()
    p_code.paragraph_format.first_line_indent = Cm(0)
    p_code.paragraph_format.line_spacing = 1.15
    p_code.paragraph_format.space_before = Pt(4)
    p_code.paragraph_format.space_after = Pt(8)
    r_code = p_code.add_run(code_text)
    r_code.bold = False
    r_code.font.name = "Consolas"
    r_code.font.size = Pt(9.5)

    # 9. Контрольні питання
    p_sec9 = doc.add_paragraph()
    p_sec9.paragraph_format.first_line_indent = Cm(1.25)
    p_sec9.paragraph_format.space_before = Pt(12)
    r = p_sec9.add_run("9. Відповіді на контрольні питання")
    r.bold = False
    
    questions = [
        ("1. Які задачі називають задачами лінійного програмування?",
         "Оптимізаційні задачі, в яких цільова функція є лінійною, а множина допустимих розв'язків задана системою лінійних рівностей або нерівностей."),
        ("2. Що таке цільова функція?",
         "Функція залежності оптимізованого критерію (прибутку, собівартості) від керованих змінних, екстремальне значення якої є метою дослідження."),
        ("3. Як записуються рівняння обмеження?",
         "У вигляді алгебраїчних нерівностей або рівнянь sum(c_ij * x_j) <= S_i, де c_ij - питомі норми витрат, а S_i - наявний запас ресурсу."),
        ("4. Які обмеження обов'язково застосовуються до задач оптимального виробництва?",
         "Обмеження за обсягами виробничих ресурсів, умови невід'ємності x_j >= 0 та вимоги цілочисельності змінних для неподільної продукції."),
        ("5. Який розв'язок ЗЛП називають оптимальним?",
         "План виробництва з області допустимих розв'язків, який надає цільовій функції екстремальне (найбільше або найменше) значення."),
        ("6. Надайте геометричну інтерпретацію ЗЛП.",
         "Знаходження точки опуклого багатокутника (ОДР), у якій лінія рівня цільової функції досягає екстремуму в напрямку вектора градієнта."),
        ("7. Яка точка допустимої множини розв'язку називається кутовою?",
         "Вершина багатокутника допустимих розв'язків, яка не може бути виражена як опукла комбінація двох інших точок цієї множини."),
        ("8. Поясніть алгоритм графічного методу розв'язання ЗЛП.",
         "Побудова прямих обмежень -> побудова ОДР -> обчислення градієнта c -> переміщення перпендикулярної лінії рівня до крайньої точки ОДР -> знаходження координат оптимуму.")
    ]
    
    for q_text, a_text in questions:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.first_line_indent = Cm(1.25)
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        p_q.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r_q = p_q.add_run(q_text + " ")
        r_q.bold = False
        r_a = p_q.add_run(a_text)
        r_a.bold = False

    # 10. Висновки
    p_conc = doc.add_paragraph()
    p_conc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_conc.paragraph_format.first_line_indent = Cm(0)
    p_conc.paragraph_format.space_before = Pt(14)
    p_conc.paragraph_format.space_after = Pt(8)
    r_conc = p_conc.add_run("Висновки")
    r_conc.bold = False
    r_conc.font.size = Pt(14)
    
    conclusions = [
        "1. У процесі виконання роботи опановано методику побудови економіко-математичних моделей та їх розв'язання графічним методом.",
        "2. За індивідуальними даними Варіанта 4 (N = 4, k = 0.5) складено модель максимізації прибутку: W = 2.0 x1 + 2.4 x2 -> max при обмеженнях на сировину (1700 од.) та машинний час (1600 год.).",
        "3. Графічним методом знайдено вершину ОДР B(161.54; 384.62), що забезпечує глобальний максимум прибутку 1246.15 грн при повному завантаженні сировини та обладнання.",
        "4. Сформовано практичний цілочисельний план випуску: 160 столів моделі А та 385 столів моделі В, що дає 1244.00 грн чистого прибутку при нульовому залишку сировини та залишку 12.0 год. машинного часу.",
        "5. Числові та графічні розрахунки верифіковано в середовищі Python за допомогою бібліотек NumPy та SciPy (метод HiGHS)."
    ]
    for c_text in conclusions:
        p_c = doc.add_paragraph()
        p_c.paragraph_format.first_line_indent = Cm(1.25)
        p_c.paragraph_format.space_after = Pt(3)
        p_c.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = p_c.add_run(c_text)
        r.bold = False

    doc.save(output_docx)
    print(f"Успішно збережено файл DOCX: {output_docx}")

if __name__ == "__main__":
    create_report_docx()
